"""Build the source-to-target mapping spec (Markdown and Excel) from
config/mapping_spec.csv and the value maps, so the two never drift apart.

    python scripts/build_mapping_spec.py
"""

from __future__ import annotations

import csv
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from recon.workbook import freeze_xlsx  # noqa: E402

SPEC = ROOT / "config" / "mapping_spec.csv"
MAPS = ROOT / "config" / "value_maps"
DOCS = ROOT / "docs"

TABLE_NAMES = {"customers": "Customers", "subscriptions": "Subscriptions", "invoices": "Invoices",
               "credit_notes": "Credit notes", "payments": "Payments"}

SCOPE = [
    ("S1", "customers", "Live after the duplicate merge (status A or H), or billed inside the window (has an in-scope invoice)."),
    ("S2", "invoices", "Invoice dated on or after 2024-10-01, or OPEN / PART at any age (DEC-BPM-002)."),
    ("S3", "subscriptions", "Customer in scope, status not PEND, and either live (ACT, TRL, SUSP, PAUS), ended on or after 2024-10-01, or billed by an in-scope invoice."),
    ("S4", "credit notes", "Follows its original invoice: in scope when that invoice is."),
    ("S5", "payments", "Follows its invoice. A payment with no matching invoice is in scope when dated on or after 2024-10-01 and the customer is in scope."),
]

VALUE_MAPS = [
    ("customer_status_map", "Customer status"),
    ("subscription_status_map", "Subscription status"),
    ("invoice_status_map", "Invoice status"),
    ("country_map", "Country"),
    ("plan_map", "Plan"),
    ("payment_method_map", "Payment method"),
]

HEADER = """# Source-to-target mapping

| | |
|---|---|
| Document ID | BPM-DW-03 |
| Programme | Billing Platform Migration (BPM), data workstream |
| Client | Northstar Systems, Inc. (fictional) |
| Version | 1.2 |
| Date | 2026-10-09 |
| Author (role) | Senior Technical Business Analyst, data workstream lead |
| Approved by | Robert Chen, Controller (financial fields); Grace Okafor, Billing Operations Manager (customer and subscription fields) |
| Status | Baselined for mock 2 |

> Generated from `config/mapping_spec.csv` and `config/value_maps/` by `scripts/build_mapping_spec.py`. Edit the CSV, not this file. The same content is in `03_source_to_target_mapping.xlsx`.

## Revision history

| Version | Date | Change |
|---|---|---|
| 1.0 | 2026-09-18 | Baselined for build 1.0 after the last mapping workshop |
| 1.1 | 2026-10-06 | M-INV-05, M-PAY-05 and M-CRN-05 reworded after DEF-BPM-001: "no time zone conversion" made explicit. No rule changed. |
| 1.2 | 2026-10-09 | M-INV-09 and M-INV-10 reworded after DEF-BPM-005 ("the platform must import it, not recalculate it"). M-CRN-09 records that credit notes carry no balance (DEC-BPM-012). No rule changed. |

## How to read this

Every target field has one row. `source` is the legacy table and column; `(platform)` means the platform generates the value and `(derived)` means it is calculated from other migrated data. `refs` points to the data quality rule (DQ-, `04_data_quality_rules.md`) that found the problem the rule handles, or the decision (DEC-, `11_decision_log.md`) that set it.

Reconciliation joins expected and actual on the legacy keys (M-*-02), never on platform ids.
"""


def _read(path: Path) -> list[dict]:
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def _md_table(rows: list[dict], cols: list[tuple[str, str]]) -> str:
    out = ["| " + " | ".join(label for _, label in cols) + " |", "|" + "---|" * len(cols)]
    for r in rows:
        out.append("| " + " | ".join(str(r.get(k, "")).replace("|", "/") for k, _ in cols) + " |")
    return "\n".join(out)


def build_markdown(spec: list[dict]) -> str:
    parts = [HEADER, "## Scope rules\n", _md_table([{"id": a, "entity": b, "rule": c} for a, b, c in SCOPE],
                                                  [("id", "Rule"), ("entity", "Entity"), ("rule", "In scope when")])]
    parts.append("\nEverything else stays in the legacy archive, which goes read-only at cutover (DEC-BPM-002).\n")
    for table, title in TABLE_NAMES.items():
        rows = [r for r in spec if r["target_table"] == table]
        parts.append(f"## {title}\n")
        parts.append(_md_table(rows, [("map_id", "Map"), ("target_field", "Target field"), ("required", "Req."),
                                      ("source", "Source"), ("rule", "Rule"), ("refs", "Refs")]))
        parts.append("")
    parts.append("## Value maps\n")
    for file, title in VALUE_MAPS:
        rows = _read(MAPS / f"{file}.csv")
        parts.append(f"### {title}\n")
        parts.append(_md_table(rows, [(k, k.replace("_", " ")) for k in rows[0].keys()]))
        parts.append("")
    return "\n".join(parts)


def build_workbook(spec: list[dict]) -> Path:
    wb = Workbook()
    head_fill, head_font = PatternFill("solid", fgColor="1F3A5F"), Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="C9CED6")
    box = Border(top=thin, bottom=thin, left=thin, right=thin)

    def sheet(ws, rows, cols):
        for j, (key, label, width) in enumerate(cols, start=1):
            c = ws.cell(row=1, column=j, value=label)
            c.font, c.fill, c.border = head_font, head_fill, box
            ws.column_dimensions[get_column_letter(j)].width = width
        for i, r in enumerate(rows, start=2):
            for j, (key, _, _) in enumerate(cols, start=1):
                c = ws.cell(row=i, column=j, value=r.get(key, ""))
                c.alignment, c.border = Alignment(wrap_text=True, vertical="top"), box
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"

    ws = wb.active
    ws.title = "Field mapping"
    sheet(ws, spec, [("map_id", "Map", 10), ("target_table", "Target table", 14), ("target_field", "Target field", 22),
                     ("required", "Req.", 6), ("source", "Source", 28), ("rule", "Rule", 70), ("refs", "Refs", 22)])
    ws = wb.create_sheet("Scope rules")
    sheet(ws, [{"id": a, "entity": b, "rule": c} for a, b, c in SCOPE],
          [("id", "Rule", 8), ("entity", "Entity", 16), ("rule", "In scope when", 100)])
    for file, title in VALUE_MAPS:
        rows = _read(MAPS / f"{file}.csv")
        ws = wb.create_sheet(title[:31])
        sheet(ws, rows, [(k, k.replace("_", " "), 22 if k != "note" else 70) for k in rows[0].keys()])
    wb.properties.creator = "Northstar BPM data workstream"
    wb.properties.title = "Source-to-target mapping v1.2"
    wb.properties.created = datetime(2026, 10, 9, 9, 0)
    path = DOCS / "03_source_to_target_mapping.xlsx"
    wb.save(path)
    freeze_xlsx(path, wb.properties.created)
    return path


def main() -> None:
    spec = _read(SPEC)
    (DOCS / "03_source_to_target_mapping.md").write_text(build_markdown(spec), encoding="utf-8", newline="\n")
    print(f"{len(spec)} field mappings -> docs/03_source_to_target_mapping.md, {build_workbook(spec).name}")


if __name__ == "__main__":
    main()
